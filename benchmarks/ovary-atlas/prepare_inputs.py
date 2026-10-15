from pathlib import Path
import csv,gzip,json,re,sys,tarfile,hashlib
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
import spatial
if len(sys.argv)!=2: raise SystemExit('Usage: prepare_inputs.py DATA_DIRECTORY')
b=Path(sys.argv[1]);d=b/'intake';d.mkdir(exist_ok=True)
meta=spatial.soft_samples(gzip.open(b/'family.soft.gz','rt').read())
wb=openpyxl.load_workbook(b/'adm7506_Data_S1.xlsx',read_only=True,data_only=True)
rows={}
for sheet in wb:
 donor=sheet.title.split()[-1];roi=None
 for idx,row in enumerate(list(sheet.values)[1:],1):
  if row[0] is not None:roi=row[0]
  rows[(donor,idx)]={'area':row[1],'nuclei':row[2],'region':row[3],'roi':roi,'cluster':row[5],'subregion':row[6] or ''}
annotation=[];allcounts={}
with tarfile.open(b/'raw.tar') as t:
 for member in sorted(t.getmembers(),key=lambda m:m.name):
  if not member.isfile():continue
  assert re.fullmatch(r'GSM\d+_DSP-[\w-]+\.dcc\.gz',member.name)
  sid=member.name.split('_')[0];title=meta[sid]['title'];match=re.fullmatch(r'Donor (\d+) AOI (\d+)',title);assert match
  donor,aoi=match.groups();r=rows[(donor,int(aoi))]
  text=gzip.decompress(t.extractfile(member).read()).decode();counts,qc=spatial.dcc(text)
  sample=member.name.split('_',1)[1][:-7];(d/(sample+'.dcc')).write_text(text)
  annotation.append({'Sample_ID':sample,'GEO_sample':sid,'donor':donor,'aoi':aoi,**r,'mapping_status':'provisional_by_donor_AOI_sequence'})
  allcounts[sample]=counts
assert len(annotation)==len(rows)==259
with (b/'sample-annotation.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=annotation[0]);w.writeheader();w.writerows(annotation)
probes=sorted(set.union(*(set(v) for v in allcounts.values())))
(b/'sparse-count-inventory.json').write_text(json.dumps({'union_probes':len(probes),'omitted_probe_entries':sum(len(probes)-len(v) for v in allcounts.values()),'omitted_entry_policy':'zero-filled for comparison; must verify against GeomxTools reader'},indent=2))
with (b/'python-imported-probes.csv').open('w',newline='') as f:
 w=csv.writer(f);w.writerow(['probe']+list(allcounts));w.writerows([p]+[v.get(p,0) for v in allcounts.values()] for p in probes)
(d/'panel.pkc').write_bytes(gzip.decompress((b/'panel.pkc.gz').read_bytes()))
refs=openpyxl.load_workbook(b/'adm7506_Data_S2.xlsx',read_only=True,data_only=True)
for sheet in refs:
 donor=sheet.title.split()[-1]
 with (b/f'published-donor{donor}-centroids.csv').open('w',newline='') as f:
  w=csv.writer(f);w.writerow(['gene']+[str(i) for i in range(1,sheet.max_column)]);w.writerows(list(sheet.values)[1:])
manifest={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [b/'raw.tar',b/'family.soft.gz',b/'panel.pkc.gz',b/'adm7506_Data_S1.xlsx',b/'adm7506_Data_S2.xlsx']}
(b/'input-hashes.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(len(probes),'probes x',len(annotation),'samples staged; provisional mapping recorded')
